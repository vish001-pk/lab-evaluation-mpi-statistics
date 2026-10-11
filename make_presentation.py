
import csv
import statistics
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "results" / "timings.csv"
GRAPHS = ROOT / "graphs"
OUTDIR = ROOT / "presentation"
OUTDIR.mkdir(exist_ok=True)

OUTPUT = OUTDIR / "Distributed_Dataset_Statistics_MPI.pptx"

# Read actual recorded trials
measurements = defaultdict(list)
with DATA.open(newline="", encoding="utf-8") as f:
    for row in csv.DictReader(f):
        key = (int(row["dataset_size"]), int(row["processes"]))
        measurements[key].append(float(row["elapsed_seconds"]))

means = {key: statistics.mean(values)
         for key, values in measurements.items()}

# Theme
NAVY = RGBColor(25, 48, 82)
BLUE = RGBColor(42, 105, 160)
TEAL = RGBColor(0, 145, 150)
PALE = RGBColor(237, 244, 249)
DARK = RGBColor(42, 52, 65)
GREY = RGBColor(100, 112, 125)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def textbox(slide, x, y, w, h, text, size=18,
            color=DARK, bold=False, font="Aptos"):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)

    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def base_slide(title, number):
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0,
        prs.slide_width, Inches(0.16)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    textbox(slide, 0.55, 0.35, 12.1, 0.55,
            title, 27, NAVY, True)

    textbox(slide, 0.58, 7.08, 10.8, 0.2,
            "Distributed Dataset Statistics using MPI",
            9, GREY)
    textbox(slide, 12.1, 7.06, 0.6, 0.25,
            str(number), 10, GREY)
    return slide


def bullets(slide, items, x=0.8, y=1.35, w=11.8,
            h=4.9, size=21):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(16)
        p.text = "\u2022  " + item
    return shape


# Slide 1: Title
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY

accent = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE,
    Inches(0.8), Inches(1.15), Inches(0.13), Inches(4.7)
)
accent.fill.solid()
accent.fill.fore_color.rgb = TEAL
accent.line.fill.background()

textbox(slide, 1.25, 1.45, 10.8, 1.55,
        "Distributed Dataset\nStatistics using MPI",
        34, WHITE, True)
textbox(slide, 1.3, 3.45, 10.5, 0.7,
        "Sequential and parallel computation of basic statistics",
        20, RGBColor(210, 224, 238))
textbox(slide, 1.3, 5.15, 10.5, 0.5,
        "PGC Review Project  |  Lab Evaluation",
        17, WHITE)
textbox(slide, 1.3, 6.6, 10, 0.35,
        "Name / Team members: add your details here",
        12, RGBColor(210, 224, 238))


# Slide 2: Problem and objectives
slide = base_slide("1. Problem Definition and Objectives", 2)
bullets(slide, [
    "Calculate the sum, average, minimum, and maximum of a dataset.",
    "Implement a sequential baseline and an MPI-based parallel version.",
    "Divide dataset indices among processes to distribute computation.",
    "Combine partial results with MPI_Reduce.",
    "Measure execution time for different dataset sizes and process counts."
], size=20)


# Slide 3: Parallel design
slide = base_slide("2. MPI Parallel Design", 3)
steps = [
    ("1", "Initialize MPI\nMPI_Init"),
    ("2", "Discover rank\nand process count"),
    ("3", "Assign each process\na portion of indices"),
    ("4", "Compute local\nsum, min, max"),
    ("5", "Combine with\nMPI_Reduce"),
    ("6", "Rank 0 computes average\nand prints results")
]
y = 1.35
for idx, (num, label) in enumerate(steps):
    col = 0 if idx < 3 else 1
    row = idx if idx < 3 else idx - 3
    x = 0.8 + col * 6.1
    yy = y + row * 1.65

    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(yy), Inches(5.4), Inches(1.15)
    )
    box.fill.solid()
    box.fill.fore_color.rgb = PALE
    box.line.color.rgb = RGBColor(206, 220, 231)

    textbox(slide, x + 0.18, yy + 0.18, 0.45, 0.5,
            num, 23, TEAL, True)
    textbox(slide, x + 0.8, yy + 0.15, 4.35, 0.8,
            label, 17, NAVY, True)

textbox(slide, 0.9, 6.55, 11.4, 0.35,
        "Only process 0 prints the global statistics.",
        16, BLUE, True)


# Slide 4: Algorithm and implementation
slide = base_slide("3. Algorithm and Implementation", 4)
textbox(slide, 0.8, 1.25, 5.5, 0.45,
        "Sequential algorithm", 21, BLUE, True)
bullets(slide, [
    "Initialize statistics.",
    "Visit each generated value.",
    "Update sum, minimum, and maximum.",
    "Compute average = sum / N."
], x=0.8, y=1.8, w=5.6, h=3.8, size=17)

textbox(slide, 6.9, 1.25, 5.5, 0.45,
        "Parallel algorithm", 21, BLUE, True)
bullets(slide, [
    "Partition indices among MPI processes.",
    "Compute local statistics.",
    "Reduce sum with MPI_SUM.",
    "Reduce minimum with MPI_MIN.",
    "Reduce maximum with MPI_MAX."
], x=6.9, y=1.8, w=5.6, h=4.3, size=17)

textbox(slide, 0.9, 6.35, 11.4, 0.45,
        "Generated values repeat from 1 to 1000: value = (index % 1000) + 1",
        15, NAVY, True)


# Slide 5: Experimental setup
slide = base_slide("4. Experimental Setup", 5)
bullets(slide, [
    "Environment: Windows, GCC, Microsoft MPI, Python, Matplotlib.",
    "Dataset sizes: 100,000; 1,000,000; and 5,000,000 values.",
    "MPI process counts: 1, 2, and 4.",
    "Five timing trials were recorded for 1M and 5M values.",
    "One timing trial was recorded for each 100K configuration.",
    "Statistics were validated against the expected generated values."
], size=19)


# Slide 6: Results table from the CSV
slide = base_slide("5. Mean MPI Runtime and Speedup", 6)
rows = sorted(means.keys())
table = slide.shapes.add_table(
    len(rows) + 1, 4,
    Inches(0.65), Inches(1.25), Inches(12.0), Inches(4.9)
).table

headers = ["Dataset size", "Processes", "Mean time (s)", "Speedup"]
for j, header in enumerate(headers):
    cell = table.cell(0, j)
    cell.text = header
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in cell.text_frame.paragraphs:
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE

table.columns[0].width = Inches(3.1)
table.columns[1].width = Inches(2.0)
table.columns[2].width = Inches(3.2)
table.columns[3].width = Inches(3.0)

for i, key in enumerate(rows, start=1):
    n, processes = key
    baseline = means.get((n, 1))
    speed = baseline / means[key] if baseline else 0
    vals = [
        f"{n:,}", str(processes),
        f"{means[key]:.6f}",
        f"{speed:.2f}x"
    ]
    for j, val in enumerate(vals):
        cell = table.cell(i, j)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE if i % 2 else PALE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.color.rgb = DARK

textbox(slide, 0.8, 6.35, 11.7, 0.45,
        "Speedup is relative to the one-process MPI run for the same dataset.",
        13, GREY)


# Slide 7: Runtime and speedup graphs
slide = base_slide("6. Performance Graphs", 7)
for filename, x, caption in [
    ("execution_time.png", 0.6, "Mean execution time"),
    ("speedup.png", 6.75, "Parallel speedup")
]:
    path = GRAPHS / filename
    if path.exists():
        slide.shapes.add_picture(
            str(path), Inches(x), Inches(1.35),
            width=Inches(5.95), height=Inches(4.65)
        )
        textbox(slide, x + 0.2, 6.15, 5.6, 0.35,
                caption, 15, NAVY, True)
    else:
        textbox(slide, x, 2.5, 5.6, 0.8,
                f"Missing graph: {filename}", 16, DARK)


# Slide 8: Efficiency and findings
slide = base_slide("7. Efficiency and Interpretation", 8)
eff = GRAPHS / "efficiency.png"
if eff.exists():
    slide.shapes.add_picture(
        str(eff), Inches(0.55), Inches(1.35),
        width=Inches(7.1), height=Inches(4.7)
    )

textbox(slide, 7.8, 1.5, 4.8, 0.45,
        "Key observations", 21, BLUE, True)
bullets(slide, [
    "For 100K values, more processes did not improve runtime.",
    "For 1M and 5M values, four processes had the lowest measured mean runtime.",
    "Small timing values and system scheduling create variability.",
    "One 100K trial per configuration limits confidence."
], x=7.8, y=2.0, w=4.9, h=4.2, size=15)

textbox(slide, 0.8, 6.35, 12.0, 0.35,
        "Results are preliminary; additional repeated trials would improve reliability.",
        13, GREY)


# Slide 9: Conclusion and demo
slide = base_slide("8. Conclusion and Demonstration", 9)
bullets(slide, [
    "The MPI program computes sum, average, minimum, and maximum correctly.",
    "The workload is distributed and local results are combined with MPI_Reduce.",
    "Larger tested datasets benefited from additional processes in the recorded runs.",
    "MPI overhead can outweigh parallel gains for small workloads.",
    "Demo: compile, then run the MPI program with 1, 2, and 4 processes."
], size=19)

textbox(slide, 0.9, 5.9, 11.8, 0.35,
        "Compile: mpicc src\\mpi_statistics.c -o mpi_statistics.exe",
        14, BLUE, True)
textbox(slide, 0.9, 6.3, 11.8, 0.35,
        "Run: mpiexec -n 4 mpi_statistics.exe 1000000",
        14, BLUE, True)

prs.save(OUTPUT)
print("PowerPoint created:", OUTPUT)
print("Slides:", len(prs.slides))
